import io

import pytest
from docx import Document
from PIL import Image

from utils.docx_utils import extract_docx, validate_docx
from utils.errors import FileValidationError

PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"


def png(size=(300, 200)):
    buffer = io.BytesIO()
    Image.new("RGB", size, (200, 200, 200)).save(buffer, format="PNG")
    return buffer.getvalue()


def make_docx(paragraphs=("Hello",), images=()):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    for image in images:
        doc.add_picture(io.BytesIO(image))
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def status_of(func, *args, **kwargs):
    with pytest.raises(FileValidationError) as info:
        func(*args, **kwargs)
    return info.value.status_code


def test_validate_accepts_docx():
    assert validate_docx("a.docx", None, make_docx()) is None


def test_validate_old_doc_is_415():
    assert status_of(validate_docx, "a.doc", None, make_docx()) == 415


def test_validate_empty_is_400():
    assert status_of(validate_docx, "a.docx", None, b"") == 400


def test_validate_garbage_is_400():
    assert status_of(validate_docx, "a.docx", None, b"PK not really a zip") == 400


def test_validate_too_large_is_413():
    assert status_of(validate_docx, "a.docx", None, b"PK" + b"\x00" * (21 * 1024 * 1024)) == 413


def test_text_comes_out_in_order():
    blocks = extract_docx(make_docx(("first", "second", "third")))["blocks"]
    assert [b["text"] for b in blocks] == ["first", "second", "third"]


def test_embedded_image_is_extracted():
    result = extract_docx(make_docx(images=[png()]))
    images = [b for b in result["blocks"] if b["type"] == "image"]
    assert len(images) == 1 and images[0]["data"].startswith(PNG_SIGNATURE)


def test_tiny_image_is_skipped():
    result = extract_docx(make_docx(images=[png((10, 10))]))
    assert result["images_skipped"] == 1
    assert not [b for b in result["blocks"] if b["type"] == "image"]


def test_max_images_truncates():
    result = extract_docx(make_docx(images=[png(), png(), png()]), max_images=2)
    assert len([b for b in result["blocks"] if b["type"] == "image"]) == 2
    assert result["images_truncated"] is True


def test_table_row_becomes_one_line():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "A", "B"
    buffer = io.BytesIO()
    doc.save(buffer)
    assert extract_docx(buffer.getvalue())["blocks"][0]["text"] == "A | B"
    